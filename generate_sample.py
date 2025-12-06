from pptx import Presentation

def create_sample_ppt():
    prs = Presentation()

    # 第一頁：標題頁
    slide_layout = prs.slide_layouts[0] # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "AI 股市新聞自動化系統"
    subtitle.text = "測試範例簡報\n使用 Streamlit 與 Python-pptx 重製"

    # 第二頁：內容頁 (痛點)
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "目前市場痛點"
    slide.placeholders[1].text = "1. 資訊爆炸：每日新聞量過大，無法人工閱讀\n2. 情緒誤判：很難量化市場恐慌指數\n3. 決策延遲：整理報告耗時，錯過交易時機"

    # 第三頁：內容頁 (解決方案)
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "AI 解決方案架構"
    slide.placeholders[1].text = "• 資料源：Google News API + Twitter\n• NLP 模型：使用 Transformer 進行摘要\n• 輸出端：自動生成每日早報 PDF"

    prs.save("sample_presentation.pptx")
    print("成功生成 sample_presentation.pptx！請將此檔案上傳至 Github。")

if __name__ == "__main__":
    create_sample_ppt()